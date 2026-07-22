from django.shortcuts import render, redirect
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib import messages
from django.db.models import Sum, Count, Q, F
from django.utils import timezone
from datetime import timedelta, datetime
from decimal import Decimal
from django.http import HttpResponse
from django.core.paginator import Paginator
import json
import jdatetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from products.models import Product
from invoices.models import Invoice, InvoiceItem
from payments.models import Payment
from customers.models import Customer
from employees.models import Employee
from products.templatetags.product_extras import to_jalali, to_jalali_datetime


def is_admin_or_employee(user):
    return user.is_authenticated and user.user_type in ['admin', 'employee']


# ============================
# داشبورد گزارشات
# ============================

@login_required
@user_passes_test(is_admin_or_employee)
def reports_dashboard(request):
    if request.user.user_type == 'employee':
        employee = request.user.employee_profile
        if not employee.access_reports:
            messages.error(request, 'شما دسترسی مشاهده گزارشات را ندارید')
            return redirect('employees:employee_panel')
    
    # ===== آمار کلی =====
    total_sales = Invoice.objects.filter(status__code='CONFIRMED').aggregate(total=Sum('total'))['total'] or 0
    total_invoices = Invoice.objects.filter(status__code='CONFIRMED').count()
    total_customers = Customer.objects.count()
    total_products_sold = InvoiceItem.objects.aggregate(total=Sum('quantity'))['total'] or 0
    
    # ===== پرفروش‌ترین محصولات =====
    top_products_data = InvoiceItem.objects.values(
        'product_code'
    ).annotate(
        total_quantity=Sum('quantity'),
        total_revenue=Sum(F('quantity') * F('price'))
    ).order_by('-total_quantity')[:10]
    
    top_products = []
    for item in top_products_data:
        product = Product.objects.filter(product_code=item['product_code']).first()
        top_products.append({
            'product_code': item['product_code'],
            'product__name': product.name if product else item['product_code'],
            'product__brand__name': product.brand.name if product and product.brand else '-',
            'total_quantity': item['total_quantity'],
            'total_revenue': item['total_revenue'],
        })
    
    # ===== پرفروش‌ترین برندها =====
    top_brands_data = InvoiceItem.objects.values(
        'product_code'
    ).annotate(
        total_quantity=Sum('quantity'),
        total_revenue=Sum(F('quantity') * F('price'))
    ).order_by('-total_revenue')[:20]
    
    brand_dict = {}
    for item in top_brands_data:
        product = Product.objects.filter(product_code=item['product_code']).first()
        brand_name = product.brand.name if product and product.brand else 'متفرقه'
        if brand_name in brand_dict:
            brand_dict[brand_name]['total_quantity'] += item['total_quantity']
            brand_dict[brand_name]['total_revenue'] += item['total_revenue']
        else:
            brand_dict[brand_name] = {
                'brand_name': brand_name,
                'total_quantity': item['total_quantity'],
                'total_revenue': item['total_revenue'],
            }
    
    top_brands = sorted(brand_dict.values(), key=lambda x: x['total_revenue'], reverse=True)[:10]
    
    # ===== مشتریان برتر =====
    top_customers = Customer.objects.annotate(
        total_purchase=Sum('invoices__total', filter=Q(invoices__status__code='CONFIRMED'))
    ).filter(
        total_purchase__isnull=False
    ).order_by('-total_purchase')[:10]
    
    # ===== داده‌های فروش ۳۰ روز اخیر =====
    end_date = timezone.now()
    start_date = end_date - timedelta(days=30)
    
    daily_sales = []
    for i in range(30):
        day = end_date - timedelta(days=i)
        day_start = day.replace(hour=0, minute=0, second=0)
        day_end = day.replace(hour=23, minute=59, second=59)
        
        sales = Invoice.objects.filter(
            date__range=(day_start, day_end),
            status__code='CONFIRMED'
        ).aggregate(total=Sum('total'))['total'] or 0
        
        daily_sales.append({
            'date': day.strftime('%Y-%m-%d'),
            'sales': float(sales)
        })
    
    context = {
        'total_sales': total_sales,
        'total_invoices': total_invoices,
        'total_customers': total_customers,
        'total_products_sold': total_products_sold,
        'top_products': top_products,
        'top_brands': top_brands,
        'top_customers': top_customers,
        'daily_sales': json.dumps(daily_sales),
    }
    return render(request, 'reports/dashboard.html', context)


# ============================
# گزارش فروش
# ============================

@login_required
@user_passes_test(is_admin_or_employee)
def sales_report(request):
    if request.user.user_type == 'employee':
        employee = request.user.employee_profile
        if not employee.access_reports:
            messages.error(request, 'شما دسترسی مشاهده گزارشات را ندارید')
            return redirect('employees:employee_panel')
    
    date_from_str = request.GET.get('date_from', '')
    date_to_str = request.GET.get('date_to', '')
    
    date_from = None
    date_to = None
    
    if date_from_str:
        try:
            parts = date_from_str.split('-')
            if len(parts) == 3:
                y, m, d = int(parts[0]), int(parts[1]), int(parts[2])
                if y < 100:
                    y = y + 1300 if y > 30 else y + 1400
                date_from = jdatetime.date(y, m, d).togregorian()
        except:
            pass
    
    if date_to_str:
        try:
            parts = date_to_str.split('-')
            if len(parts) == 3:
                y, m, d = int(parts[0]), int(parts[1]), int(parts[2])
                if y < 100:
                    y = y + 1300 if y > 30 else y + 1400
                date_to = jdatetime.date(y, m, d).togregorian()
        except:
            pass
    
    invoices = Invoice.objects.filter(status__code='CONFIRMED')
    
    if date_from:
        invoices = invoices.filter(date__gte=date_from)
    if date_to:
        invoices = invoices.filter(date__lte=date_to)
    
    total_sales = invoices.aggregate(total=Sum('total'))['total'] or 0
    total_discount = invoices.aggregate(total=Sum('discount'))['total'] or 0
    invoice_count = invoices.count()
    
    sales_by_day = invoices.extra(
        {'day': "date(invoices_invoice.date)"}
    ).values('day').annotate(
        total=Sum('total'),
        count=Count('id')
    ).order_by('-day')[:30]
    
    sales_by_customer = invoices.values(
        'customer__user__first_name',
        'customer__user__last_name'
    ).annotate(
        total=Sum('total'),
        count=Count('id')
    ).order_by('-total')[:20]
    
    context = {
        'total_sales': total_sales,
        'total_discount': total_discount,
        'invoice_count': invoice_count,
        'sales_by_day': sales_by_day,
        'sales_by_customer': sales_by_customer,
        'date_from': date_from_str,
        'date_to': date_to_str,
    }
    return render(request, 'reports/sales_report.html', context)


# ============================
# گزارش پرداخت‌ها
# ============================

@login_required
@user_passes_test(is_admin_or_employee)
def payments_report(request):
    if request.user.user_type == 'employee':
        employee = request.user.employee_profile
        if not employee.access_reports:
            messages.error(request, 'شما دسترسی مشاهده گزارشات را ندارید')
            return redirect('employees:employee_panel')
    
    date_from_str = request.GET.get('date_from', '')
    date_to_str = request.GET.get('date_to', '')
    
    date_from = None
    date_to = None
    
    if date_from_str:
        try:
            parts = date_from_str.split('-')
            if len(parts) == 3:
                y, m, d = int(parts[0]), int(parts[1]), int(parts[2])
                if y < 100:
                    y = y + 1300 if y > 30 else y + 1400
                date_from = jdatetime.date(y, m, d).togregorian()
        except:
            pass
    
    if date_to_str:
        try:
            parts = date_to_str.split('-')
            if len(parts) == 3:
                y, m, d = int(parts[0]), int(parts[1]), int(parts[2])
                if y < 100:
                    y = y + 1300 if y > 30 else y + 1400
                date_to = jdatetime.date(y, m, d).togregorian()
        except:
            pass
    
    payments = Payment.objects.filter(is_confirmed=True, status='confirmed')
    
    if date_from:
        payments = payments.filter(date__gte=date_from)
    if date_to:
        payments = payments.filter(date__lte=date_to)
    
    total_payments = payments.aggregate(total=Sum('amount'))['total'] or 0
    payment_count = payments.count()
    
    payments_by_type = payments.values('payment_type__name').annotate(
        total=Sum('amount'),
        count=Count('id')
    ).order_by('-total')
    
    payments_by_customer = payments.values(
        'customer__user__first_name',
        'customer__user__last_name'
    ).annotate(
        total=Sum('amount'),
        count=Count('id')
    ).order_by('-total')[:20]
    
    context = {
        'total_payments': total_payments,
        'payment_count': payment_count,
        'payments_by_type': payments_by_type,
        'payments_by_customer': payments_by_customer,
        'date_from': date_from_str,
        'date_to': date_to_str,
    }
    return render(request, 'reports/payments_report.html', context)


# ============================
# گزارش محصولات
# ============================

@login_required
@user_passes_test(is_admin_or_employee)
def products_report(request):
    if request.user.user_type == 'employee':
        employee = request.user.employee_profile
        if not employee.access_reports:
            messages.error(request, 'شما دسترسی مشاهده گزارشات را ندارید')
            return redirect('employees:employee_panel')
    
    date_from_str = request.GET.get('date_from', '')
    date_to_str = request.GET.get('date_to', '')
    
    date_from = None
    date_to = None
    
    if date_from_str:
        try:
            parts = date_from_str.split('-')
            if len(parts) == 3:
                y, m, d = int(parts[0]), int(parts[1]), int(parts[2])
                if y < 100:
                    y = y + 1300 if y > 30 else y + 1400
                date_from = jdatetime.date(y, m, d).togregorian()
        except:
            pass
    
    if date_to_str:
        try:
            parts = date_to_str.split('-')
            if len(parts) == 3:
                y, m, d = int(parts[0]), int(parts[1]), int(parts[2])
                if y < 100:
                    y = y + 1300 if y > 30 else y + 1400
                date_to = jdatetime.date(y, m, d).togregorian()
        except:
            pass
    
    items = InvoiceItem.objects.filter(invoice__status__code='CONFIRMED')
    
    if date_from:
        items = items.filter(invoice__date__gte=date_from)
    if date_to:
        items = items.filter(invoice__date__lte=date_to)
    
    top_products_data = items.values('product_code').annotate(
        total_quantity=Sum('quantity'),
        total_revenue=Sum(F('quantity') * F('price')),
        total_points=Sum('points_earned')
    ).order_by('-total_quantity')[:50]
    
    top_products = []
    for item in top_products_data:
        product = Product.objects.filter(product_code=item['product_code']).first()
        top_products.append({
            'product_code': item['product_code'],
            'product__name': product.name if product else item['product_code'],
            'product__brand__name': product.brand.name if product and product.brand else '-',
            'total_quantity': item['total_quantity'],
            'total_revenue': item['total_revenue'],
            'total_points': item['total_points'],
        })
    
    total_products_sold = items.aggregate(total=Sum('quantity'))['total'] or 0
    total_revenue = items.aggregate(total=Sum(F('quantity') * F('price')))['total'] or 0
    unique_products = items.values('product_code').distinct().count()
    
    context = {
        'top_products': top_products,
        'total_products_sold': total_products_sold,
        'total_revenue': total_revenue,
        'unique_products': unique_products,
        'date_from': date_from_str,
        'date_to': date_to_str,
    }
    return render(request, 'reports/products_report.html', context)


# ============================
# گزارش مشتریان
# ============================

@login_required
@user_passes_test(is_admin_or_employee)
def customers_report(request):
    if request.user.user_type == 'employee':
        employee = request.user.employee_profile
        if not employee.access_reports:
            messages.error(request, 'شما دسترسی مشاهده گزارشات را ندارید')
            return redirect('employees:employee_panel')
    
    customers = Customer.objects.annotate(
        total_purchase=Sum('invoices__total', filter=Q(invoices__status__code='CONFIRMED')),
        total_payment=Sum('payments__amount', filter=Q(payments__is_confirmed=True)),
        invoice_count=Count('invoices', filter=Q(invoices__status__code='CONFIRMED'))
    ).order_by('-total_purchase')
    
    debt_filter = request.GET.get('debt', '')
    if debt_filter == 'has_debt':
        customers = customers.filter(total_purchase__gt=F('total_payment'))
    elif debt_filter == 'no_debt':
        customers = customers.filter(total_purchase__lte=F('total_payment'))
    
    search = request.GET.get('search', '')
    if search:
        customers = customers.filter(
            Q(user__first_name__icontains=search) |
            Q(user__last_name__icontains=search) |
            Q(user__national_id__icontains=search)
        )
    
    paginator = Paginator(customers, 30)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)
    
    total_customers = customers.count()
    total_debt = 0
    total_credit = 0
    
    for c in customers:
        debt = c.total_purchase - c.total_payment
        if debt > 0:
            total_debt += debt
        else:
            total_credit += abs(debt)
    
    context = {
        'customers': page_obj,
        'total_customers': total_customers,
        'total_debt': total_debt,
        'total_credit': total_credit,
        'search': search,
        'debt_filter': debt_filter,
    }
    return render(request, 'reports/customers_report.html', context)


# ============================
# خروجی اکسل
# ============================

@login_required
@user_passes_test(is_admin_or_employee)
def export_to_excel(request, report_type):
    if request.user.user_type == 'employee':
        employee = request.user.employee_profile
        if not employee.access_reports:
            messages.error(request, 'شما دسترسی خروجی اکسل را ندارید')
            return redirect('employees:employee_panel')
    
    date_from_str = request.GET.get('date_from', '')
    date_to_str = request.GET.get('date_to', '')
    
    date_from = None
    date_to = None
    
    if date_from_str:
        try:
            parts = date_from_str.split('-')
            if len(parts) == 3:
                y, m, d = int(parts[0]), int(parts[1]), int(parts[2])
                if y < 100:
                    y = y + 1300 if y > 30 else y + 1400
                date_from = jdatetime.date(y, m, d).togregorian()
        except:
            pass
    
    if date_to_str:
        try:
            parts = date_to_str.split('-')
            if len(parts) == 3:
                y, m, d = int(parts[0]), int(parts[1]), int(parts[2])
                if y < 100:
                    y = y + 1300 if y > 30 else y + 1400
                date_to = jdatetime.date(y, m, d).togregorian()
        except:
            pass
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    
    header_font = Font(name='B Nazanin', size=12, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    filename = f"report_{report_type}_{date_from_str}_{date_to_str}.xlsx"
    
    # ===== گزارش فروش =====
    if report_type == 'sales':
        invoices = Invoice.objects.filter(status__code='CONFIRMED')
        if date_from:
            invoices = invoices.filter(date__gte=date_from)
        if date_to:
            invoices = invoices.filter(date__lte=date_to)
        
        headers = ['ردیف', 'شماره فاکتور', 'نام مشتری', 'تاریخ', 'مبلغ کل', 'تخفیف', 'مبلغ نهایی']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
        
        for idx, invoice in enumerate(invoices, 1):
            ws.cell(row=idx+1, column=1, value=idx).border = border
            ws.cell(row=idx+1, column=2, value=invoice.invoice_id).border = border
            ws.cell(row=idx+1, column=3, value=f"{invoice.customer.user.first_name} {invoice.customer.user.last_name}").border = border
            ws.cell(row=idx+1, column=4, value=to_jalali(invoice.date)).border = border
            ws.cell(row=idx+1, column=5, value=float(invoice.subtotal)).border = border
            ws.cell(row=idx+1, column=6, value=float(invoice.discount)).border = border
            ws.cell(row=idx+1, column=7, value=float(invoice.total)).border = border
        
        column_widths = [8, 20, 25, 15, 15, 15, 15]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
    
    # ===== گزارش پرداخت‌ها =====
    elif report_type == 'payments':
        payments = Payment.objects.filter(is_confirmed=True, status='confirmed')
        if date_from:
            payments = payments.filter(date__gte=date_from)
        if date_to:
            payments = payments.filter(date__lte=date_to)
        
        headers = ['ردیف', 'شماره پرداخت', 'نام مشتری', 'تاریخ', 'مبلغ', 'نوع پرداخت', 'کد پیگیری', 'وضعیت']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
        
        for idx, payment in enumerate(payments, 1):
            ws.cell(row=idx+1, column=1, value=idx).border = border
            ws.cell(row=idx+1, column=2, value=payment.payment_id).border = border
            ws.cell(row=idx+1, column=3, value=f"{payment.customer.user.first_name} {payment.customer.user.last_name}").border = border
            ws.cell(row=idx+1, column=4, value=to_jalali_datetime(payment.date)).border = border
            ws.cell(row=idx+1, column=5, value=float(payment.amount)).border = border
            ws.cell(row=idx+1, column=6, value=payment.payment_type.name).border = border
            ws.cell(row=idx+1, column=7, value=payment.confirmation_code or '-').border = border
            ws.cell(row=idx+1, column=8, value='تایید شده' if payment.is_confirmed else 'در انتظار تایید').border = border
        
        column_widths = [8, 20, 25, 20, 15, 15, 20, 15]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
    
    # ===== گزارش محصولات =====
    elif report_type == 'products':
        items = InvoiceItem.objects.filter(invoice__status__code='CONFIRMED')
        if date_from:
            items = items.filter(invoice__date__gte=date_from)
        if date_to:
            items = items.filter(invoice__date__lte=date_to)
        
        products_data = items.values('product_code').annotate(
            total_quantity=Sum('quantity'),
            total_revenue=Sum(F('quantity') * F('price')),
            total_points=Sum('points_earned')
        ).order_by('-total_quantity')
        
        headers = ['ردیف', 'کد محصول', 'نام محصول', 'برند', 'تعداد فروش', 'درآمد کل', 'امتیاز']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
        
        for idx, product in enumerate(products_data, 1):
            p = Product.objects.filter(product_code=product['product_code']).first()
            ws.cell(row=idx+1, column=1, value=idx).border = border
            ws.cell(row=idx+1, column=2, value=product['product_code'] or '-').border = border
            ws.cell(row=idx+1, column=3, value=p.name if p else product['product_code']).border = border
            ws.cell(row=idx+1, column=4, value=p.brand.name if p and p.brand else '-').border = border
            ws.cell(row=idx+1, column=5, value=product['total_quantity']).border = border
            ws.cell(row=idx+1, column=6, value=float(product['total_revenue'])).border = border
            ws.cell(row=idx+1, column=7, value=product['total_points']).border = border
        
        column_widths = [8, 15, 40, 20, 15, 20, 15]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
    
    # ===== گزارش مشتریان =====
    elif report_type == 'customers':
        customers = Customer.objects.annotate(
            total_purchase=Sum('invoices__total', filter=Q(invoices__status__code='CONFIRMED')),
            total_payment=Sum('payments__amount', filter=Q(payments__is_confirmed=True)),
            invoice_count=Count('invoices', filter=Q(invoices__status__code='CONFIRMED'))
        ).order_by('-total_purchase')
        
        headers = ['ردیف', 'نام مشتری', 'کد ملی', 'تلفن', 'اعتبار', 'بدهی', 'کل خرید', 'تعداد فاکتور', 'امتیاز']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
        
        for idx, customer in enumerate(customers, 1):
            debt = customer.total_purchase - customer.total_payment
            ws.cell(row=idx+1, column=1, value=idx).border = border
            ws.cell(row=idx+1, column=2, value=f"{customer.user.first_name} {customer.user.last_name}").border = border
            ws.cell(row=idx+1, column=3, value=customer.user.national_id).border = border
            ws.cell(row=idx+1, column=4, value=customer.user.phone).border = border
            ws.cell(row=idx+1, column=5, value=float(customer.credit or 0)).border = border
            ws.cell(row=idx+1, column=6, value=float(debt if debt > 0 else 0)).border = border
            ws.cell(row=idx+1, column=7, value=float(customer.total_purchase or 0)).border = border
            ws.cell(row=idx+1, column=8, value=customer.invoice_count or 0).border = border
            ws.cell(row=idx+1, column=9, value=customer.total_points or 0).border = border
        
        column_widths = [8, 25, 15, 15, 15, 15, 20, 15, 15]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
    
    else:
        filename = "report.xlsx"
    
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    
    return response